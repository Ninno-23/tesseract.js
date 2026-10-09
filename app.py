from __future__ import annotations
import io, json, os, re, csv, zipfile
from pathlib import Path
from typing import Any

import requests
from flask import Flask, Response, jsonify, request, send_from_directory

BASE_DIR = Path(__file__).resolve().parent
app = Flask(__name__, static_folder="static", static_url_path="/static")
app.config["MAX_CONTENT_LENGTH"] = int(os.getenv("MAX_UPLOAD_MB", "20")) * 1024 * 1024

DEFAULT_CHAT_MODEL = os.getenv("AI_MODEL", "openai/gpt-oss-120b:fastest")
DEFAULT_IMAGE_MODEL = os.getenv("IMAGE_MODEL", "black-forest-labs/FLUX.1-dev")
CHAT_URL = os.getenv("AI_API_URL", "https://router.huggingface.co/v1/chat/completions")
IMAGE_URL = os.getenv("IMAGE_API_URL", "https://router.huggingface.co/hf-inference/models/{model}")
SYSTEM_PROMPT = ("You are Carrot AI, a capable, accurate and helpful assistant. Answer directly, "
                 "use supplied study material carefully, do not invent facts, and clearly say when uncertain.")


def chat_model() -> str:
    candidate = request.headers.get("X-Chat-Model", "").strip()
    return candidate[:180] if candidate and re.fullmatch(r"[A-Za-z0-9_.:/-]+", candidate) else DEFAULT_CHAT_MODEL


def image_model() -> str:
    candidate = request.headers.get("X-Image-Model", "").strip()
    return candidate[:180] if candidate and re.fullmatch(r"[A-Za-z0-9_.:/-]+", candidate) else DEFAULT_IMAGE_MODEL


def response_style() -> str:
    style = request.headers.get("X-Response-Style", "balanced").lower()
    return style if style in {"balanced", "concise", "detailed", "tutor"} else "balanced"


def api_key(image: bool = False) -> str:
    if image:
        return os.getenv("IMAGE_API_KEY", "").strip() or os.getenv("AI_API_KEY", "").strip()
    return os.getenv("AI_API_KEY", "").strip()


def provider_error(resp: requests.Response) -> str:
    try:
        data = resp.json()
        err = data.get("error", data)
        if isinstance(err, dict):
            err = err.get("message") or err.get("detail") or str(err)
        return str(err)[:500]
    except Exception:
        return (resp.text or f"Provider returned HTTP {resp.status_code}")[:500]


def call_chat(messages: list[dict[str, str]], max_tokens: int = 1800) -> str:
    key = api_key()
    if not key:
        raise RuntimeError("AI provider is not configured. Add AI_API_KEY to your hosting service's private environment variables, then restart the app.")
    style = response_style()
    style_hint = {
        "concise": "Keep the answer concise while still being accurate.",
        "detailed": "Give a detailed, well-structured explanation with examples when useful.",
        "tutor": "Act as a patient study tutor: explain step by step, check understanding, and define difficult terms.",
        "balanced": "Use a clear, helpful level of detail."
    }[style]
    prepared = list(messages)
    if prepared and prepared[0].get("role") == "system":
        prepared[0] = {"role": "system", "content": prepared[0].get("content", SYSTEM_PROMPT) + "\n" + style_hint}
    else:
        prepared.insert(0, {"role": "system", "content": SYSTEM_PROMPT + "\n" + style_hint})
    try:
        resp = requests.post(CHAT_URL, headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
            json={"model": chat_model(), "messages": prepared, "temperature": 0.35, "max_tokens": max_tokens}, timeout=(10, 120))
    except requests.RequestException as exc:
        raise RuntimeError(f"Could not reach the AI provider: {exc.__class__.__name__}. Check the server network and AI_API_URL.") from exc
    if not resp.ok:
        raise RuntimeError(f"AI provider returned HTTP {resp.status_code}: {provider_error(resp)}")
    try:
        data = resp.json()
        content = data["choices"][0]["message"]["content"]
        if isinstance(content, list):
            content = "".join(str(part.get("text", "")) for part in content if isinstance(part, dict))
        content = str(content).strip()
    except (ValueError, KeyError, IndexError, TypeError) as exc:
        raise RuntimeError("The AI provider returned an unexpected response format.") from exc
    if not content:
        raise RuntimeError("The AI provider returned an empty answer. Check the model and provider quota.")
    return content


def parse_json_answer(text: str) -> dict[str, Any]:
    clean = re.sub(r"^\s*```(?:json)?\s*|\s*```\s*$", "", text, flags=re.I)
    try:
        return json.loads(clean)
    except json.JSONDecodeError:
        start, end = clean.find("{"), clean.rfind("}")
        if start >= 0 and end > start:
            return json.loads(clean[start:end + 1])
        raise ValueError("The AI returned malformed structured content. Please try again.")


def extract_upload() -> str:
    text = request.form.get("text", "").strip()
    if text:
        return text[:90000]
    file = request.files.get("file")
    if not file:
        raise ValueError("Paste text or upload a document first.")
    name = file.filename or "upload"
    ext = Path(name).suffix.lower()
    raw = file.read()
    if len(raw) > 20 * 1024 * 1024:
        raise ValueError("File exceeds the 20 MB limit.")
    if ext in {".txt", ".md", ".py", ".js", ".ts", ".json", ".csv", ".html", ".css", ".sql", ".yaml", ".yml"}:
        body = raw.decode("utf-8", errors="replace")
    elif ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            pages = []
            for i, page in enumerate(reader.pages[:100], 1):
                pages.append(f"[Page {i}]\n{page.extract_text() or ''}")
            body = "\n\n".join(pages).strip()
            if not body:
                raise ValueError("This PDF appears to be scanned or image-only. Use the browser OCR flow or upload page images.")
        except ImportError as exc:
            raise ValueError("PDF reader dependency is unavailable. Reinstall requirements.txt.") from exc
    elif ext == ".docx":
        try:
            from docx import Document
            doc = Document(io.BytesIO(raw))
            body = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                body += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        except ImportError as exc:
            raise ValueError("Word reader dependency is unavailable. Reinstall requirements.txt.") from exc
    elif ext in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sections = []
            for sheet in workbook.worksheets:
                lines = [f"SHEET: {sheet.title}"]
                for row in sheet.iter_rows(values_only=True):
                    vals = ["" if v is None else str(v) for v in row]
                    if any(vals): lines.append(" | ".join(vals))
                sections.append("\n".join(lines))
            body = "\n\n".join(sections)
        except ImportError as exc:
            raise ValueError("Spreadsheet reader dependency is unavailable. Reinstall requirements.txt.") from exc
    elif ext == ".zip":
        parts = []
        allowed = {".txt", ".md", ".csv", ".json", ".html", ".css", ".js", ".ts", ".py", ".sql", ".yaml", ".yml"}
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                for info in archive.infolist():
                    if info.is_dir() or Path(info.filename).suffix.lower() not in allowed: continue
                    if sum(len(x) for x in parts) > 80000: break
                    parts.append(f"FILE: {info.filename}\n" + archive.read(info).decode("utf-8", errors="replace")[:12000])
            body = "\n\n".join(parts)
        except zipfile.BadZipFile as exc:
            raise ValueError("The uploaded ZIP archive is invalid.") from exc
        if not body: raise ValueError("No supported text or code files were found in this ZIP.")
    elif ext in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        raise ValueError("Image OCR runs in the browser. Please use the OCR scanner to extract text, then submit the extracted text for AI analysis.")
    else:
        raise ValueError("Unsupported file type. Use PDF, DOCX, XLSX/XLSM, ZIP with text/code files, TXT, Markdown, images, or supported code/data files.")
    body = body.strip()
    if not body: raise ValueError("No readable text was found in this document.")
    return f"Document: {name}\n\n{body[:90000]}"


@app.get("/")
def index():
    return send_from_directory(BASE_DIR, "index.html")


@app.get("/api/health")
def health():
    key = bool(api_key())
    image_key = bool(api_key(image=True))
    return jsonify(ok=True, version="1.0.0-fullstack", mode="server", ai_configured=key,
                   image_configured=image_key, model=chat_model(), image_model=image_model(),
                   settings_mode="server-environment", max_upload_mb=app.config["MAX_CONTENT_LENGTH"] // (1024 * 1024))


@app.post("/api/chat/stream")
def chat_stream():
    data = request.get_json(silent=True) or {}
    message = str(data.get("message", "")).strip()
    if not message: return jsonify(error="Enter a message first."), 400
    history = data.get("history", [])
    messages = [{"role": "system", "content": SYSTEM_PROMPT}]
    if isinstance(history, list):
        for item in history[-16:]:
            if isinstance(item, dict) and item.get("role") in {"user", "assistant"}:
                messages.append({"role": item["role"], "content": str(item.get("content", ""))[:9000]})
    messages.append({"role": "user", "content": message[:90000]})
    try:
        answer = call_chat(messages, 2200)
        payload = "data: " + json.dumps({"text": answer}, ensure_ascii=False) + "\n\ndata: [DONE]\n\n"
        return Response(payload, mimetype="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})
    except Exception as exc:
        return jsonify(error=str(exc)), 503


@app.post("/api/study/<kind>")
def study(kind: str):
    if kind not in {"summarize", "flashcards", "quiz"}: return jsonify(error="Unknown study tool."), 404
    try:
        src = extract_upload()
        if kind == "summarize":
            prompt = "Create a grounded study summary: overview, key ideas, definitions/names/dates/formulas, and a short self-check. Use only this source; do not invent facts.\n\n" + src
        elif kind == "flashcards":
            prompt = 'Create 12 to 24 accurate study flashcards using only the source. Return ONLY JSON: {"cards":[{"front":"question or term","back":"answer"}]}. Fewer are fine for short sources.\n\n' + src
        else:
            prompt = 'Create 8 to 12 multiple-choice questions using only the source. Return ONLY JSON: {"questions":[{"question":"...","options":["A","B","C","D"],"answer":0,"explanation":"..."}]}. Answer is index 0-3.\n\n' + src
        answer = call_chat([{"role": "system", "content": SYSTEM_PROMPT}, {"role": "user", "content": prompt}], 3000)
        if kind == "summarize": return jsonify(answer=answer)
        obj = parse_json_answer(answer)
        if kind == "flashcards":
            cards = []
            for item in obj.get("cards", obj.get("flashcards", [])):
                if not isinstance(item, dict): continue
                front, back = str(item.get("front", item.get("q", ""))).strip(), str(item.get("back", item.get("a", ""))).strip()
                if front and back: cards.append({"front": front[:1000], "back": back[:2000]})
                if len(cards) >= 40: break
            if not cards: raise ValueError("No valid flashcards returned. Try a longer source.")
            return jsonify(answer=f"Created {len(cards)} flashcards.", structured=cards)
        questions = []
        for item in obj.get("questions", obj.get("quiz", [])):
            if not isinstance(item, dict): continue
            options, answer_idx = item.get("options"), item.get("answer")
            if item.get("question") and isinstance(options, list) and len(options) == 4 and str(answer_idx).isdigit() and 0 <= int(answer_idx) <= 3:
                questions.append({"question": str(item["question"]), "options": [str(x) for x in options], "answer": int(answer_idx), "explanation": str(item.get("explanation", ""))})
            if len(questions) >= 20: break
        if not questions: raise ValueError("No valid quiz questions returned. Try a longer source.")
        return jsonify(answer=f"Created a {len(questions)}-question quiz.", structured=questions)
    except Exception as exc:
        return jsonify(error=str(exc)), 400


@app.post("/api/document")
def document():
    try:
        src = extract_upload()
        answer = call_chat([{"role": "system", "content": SYSTEM_PROMPT + " Analyze uploaded documents carefully. Preserve key facts and flag uncertainty."},
                            {"role": "user", "content": "Analyze this document: overview, key points, important terms/numbers, and next steps. Do not invent details.\n\n" + src}], 2200)
        return jsonify(answer=answer)
    except Exception as exc:
        return jsonify(error=str(exc)), 400


@app.post("/api/image")
def image_generation():
    key = api_key(image=True)
    if not key: return jsonify(error="Image generation is not configured. Set IMAGE_API_KEY or AI_API_KEY on the server."), 503
    data = request.get_json(silent=True) or {}
    prompt = str(data.get("prompt", "")).strip()
    if not prompt: return jsonify(error="Describe the image you want to create."), 400
    model = image_model()
    url = IMAGE_URL.replace("{model}", model)
    try:
        resp = requests.post(url, headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json", "Accept": "image/*"},
                             json={"inputs": prompt[:3000], "options": {"wait_for_model": True}}, timeout=(10, 180))
    except requests.RequestException as exc:
        return jsonify(error=f"Could not reach the image provider: {exc.__class__.__name__}. Check IMAGE_API_URL."), 503
    if not resp.ok: return jsonify(error=f"Image provider returned HTTP {resp.status_code}: {provider_error(resp)}"), 503
    content_type = resp.headers.get("Content-Type", "application/octet-stream").split(";", 1)[0]
    if not content_type.startswith("image/"):
        return jsonify(error="Image provider returned non-image data. Check IMAGE_API_URL and IMAGE_MODEL."), 502
    return Response(resp.content, content_type=content_type, headers={"Cache-Control": "no-store", "X-Content-Type-Options": "nosniff"})


@app.post("/api/search")
def search():
    data = request.get_json(silent=True) or {}
    query = str(data.get("query", "")).strip()
    if not query: return jsonify(error="Enter a search query first."), 400
    from urllib.parse import quote_plus
    url = "https://duckduckgo.com/?q=" + quote_plus(query)
    return jsonify(query=query, results=[{"title": "Open web search results", "url": url, "snippet": "Open the results to review current sources."}], synthesis="Search results are opened externally. Verify sources before relying on them.")


@app.errorhandler(413)
def too_large(_):
    return jsonify(error=f"Upload too large. Maximum size is {app.config['MAX_CONTENT_LENGTH'] // (1024 * 1024)} MB."), 413

@app.after_request
def security_headers(resp):
    resp.headers.setdefault("X-Content-Type-Options", "nosniff")
    resp.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    resp.headers.setdefault("X-Frame-Options", "DENY")
    if request.path.startswith("/api/"):
        resp.headers.setdefault("Cache-Control", "no-store")
    return resp

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5000")))
